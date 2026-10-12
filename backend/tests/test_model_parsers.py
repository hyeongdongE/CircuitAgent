"""실제 공개 수치와 원문/위상을 검증하고 오류 입력은 별도로 작성한다."""

import hashlib
import io
import json
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest
import xlwt
from circuitagent.bom import MAX_BYTES, csv_cells, parse_bom
from circuitagent.conversion import convert
from circuitagent.domain import CircuitDesign, PinKey, wire
from circuitagent.graph import CircuitGraph, assemble
from circuitagent.netlist import KiCadNetlistParser, QcvNetlistParser, sexpression
from circuitagent.units import normalize
from conftest import DATA, ROOT
from pydantic import ValidationError


@pytest.mark.parametrize(
    "raw,unit,value",
    [
        ("1kΩ", "Ω", "1000"),
        ("4R7", "Ω", "4.7"),
        ("0R", "Ω", "0"),
        ("1M", "Ω", "1000000"),
        ("1m", "Ω", ".001"),
        ("100nF", "F", ".0000001"),
        ("4u7", "F", ".0000047"),
        ("2,2u", "F", ".0000022"),
        ("1uH", "H", ".000001"),
        ("6V3", "V", "6.3"),
        ("1A5", "A", "1.5"),
        ("±1%", "%", "1"),
        ("0", "F", "0"),
        ("1e3Ω", "Ω", "1000"),
        ("47K", "Ω", "47000"),
        ("100µF", "F", ".0001"),
        ("100μF", "F", ".0001"),
        (".5A", "A", ".5"),
    ],
)
def test_units(raw, unit, value):
    a = normalize(raw, unit)
    assert a.normalized == Decimal(value) and a.raw == raw and a.unit == unit


@pytest.mark.parametrize(
    "raw,unit",
    [
        ("1,000", "Ω"),
        ("5V", "Ω"),
        ("1R5", "F"),
        ("~", "Ω"),
        ("NaN", "Ω"),
        ("inf", "F"),
        ("1e999", "Ω"),
        ("2k%", "%"),
    ],
)
def test_bad_units(raw, unit):
    with pytest.raises(ValueError):
        normalize(raw, unit)


@pytest.mark.parametrize(
    "board,count,nets,connected,separate,term,dnp,issues",
    [
        ("fe", 445, 335, 1500, 143, 135, 51, 80),
        ("ntop", 296, 233, 972, 147, 139, 67, 93),
        ("fe-ntop", 161, 193, 794, 25, 17, 37, 129),
    ],
)
def test_real_boards(boards, board, count, nets, connected, separate, term, dnp, issues):
    b, n, d, g = boards[board]
    assert len(b.components) == count and len(n.components) == count + 8
    assert set(b.components) <= set(n.components)
    assert len(n.nets) == nets and len(n.connections) == connected
    assert int(n.source_metadata["separatePinRecords"]) == separate
    assert sum(p.disposition == "TERMINATED" for p in n.pins) == term
    assert sum(c.assembly_status == "DNP" for c in b.components.values()) == dnp
    assert len(d.issues) == issues
    assert not any(i.code == "INVALID_UNIT" for i in d.issues)


def test_actual_sources_checksums():
    for r in json.loads((ROOT / "scripts/qcv-fixtures.json").read_text())["files"]:
        b = (DATA / r["local"]).read_bytes()
        assert len(b) == r["bytes"] and hashlib.sha256(b).hexdigest() == r["sha256"]


@pytest.mark.parametrize("board", ["fe", "ntop", "fe-ntop"])
def test_xlsx_conversion(tmp_path, board):
    for b in ["fe", "ntop", "fe-ntop"]:
        (tmp_path / (b + ".csv")).write_bytes((DATA / (b + ".csv")).read_bytes())
    convert(tmp_path)
    expected = csv_cells((tmp_path / (board + ".csv")).read_bytes())
    w = openpyxl.load_workbook(tmp_path / (board + ".xlsx"))
    for i, row in enumerate(expected, 1):
        assert [w.active.cell(i, j).value or "" for j in range(1, len(row) + 1)] == row
    w.close()


@pytest.mark.parametrize("name,count,total", [("5v", 12, 12), ("shift", 27, 31)])
def test_kicad(name, count, total):
    p = Path(__file__).parent / "fixtures" / name
    b = parse_bom(next(p.glob("*.csv")).read_bytes(), "bom.csv")
    n = KiCadNetlistParser().parse(next(p.glob("*.net")).read_bytes(), "net.net")
    d = assemble(name, "v1", b, n)
    g = CircuitGraph(d)
    assert len(b.components) == count and len(d.components) == total
    if name == "5v":
        assert (
            len(d.pins) == 25
            and len(d.nets) == 6
            and len(d.connections) == 25
            and g.get_component_attributes("R1")["Resistance"].normalized == 1000
            and g.get_component("U1").value == "L7805"
        )
    else:
        assert {"H1", "H2", "H3", "H4"} <= set(d.components)


@pytest.mark.parametrize(
    "flag,expected",
    [
        ("yes", "DNP"),
        ("no", "FITTED"),
        ("", "UNSPECIFIED"),
        ("maybe", "UNSPECIFIED"),
        ("1", "DNP"),
        ("0", "FITTED"),
    ],
)
def test_dnp(flag, expected):
    b = parse_bom(f"Reference,Value,DNP\nR1,0,{flag}\n".encode(), "bom.csv")
    assert b.components["R1"].assembly_status == expected
    assert next(iter(b.specifications.values())).attributes["Resistance"].normalized == 0


@pytest.mark.parametrize("kind", ["xlsx", "xls"])
def test_excel(kind):
    rows = [["MPN", "Value", "Reference", "Tolerance"], ["", "1k", "R1", ""]]
    out = io.BytesIO()
    if kind == "xlsx":
        w = openpyxl.Workbook()
        for row in rows:
            w.active.append(row)
        w.save(out)
        w.close()
    else:
        w = xlwt.Workbook()
        s = w.add_sheet("BOM")
        for i, row in enumerate(rows):
            for j, v in enumerate(row):
                s.write(i, j, v)
        w.save(out)
    b = parse_bom(out.getvalue(), "bom." + kind)
    spec = next(iter(b.specifications.values()))
    assert (
        spec.mpn is None
        and "Tolerance" not in spec.attributes
        and b.components["R1"].raw_sources["bom"]["MPN"] == ""
    )


@pytest.mark.parametrize(
    "data,name",
    [
        (b"", "bom.csv"),
        (b"Value\n1k", "bom.csv"),
        (b"Reference,Ref\nR1,R2", "bom.csv"),
        (b"Reference\nR1,extra", "bom.csv"),
        (b"bad", "bom.xlsx"),
        (b"bad", "bom.xls"),
        (b"\xff", "bom.csv"),
        (b"bad", "bom.txt"),
    ],
)
def test_bad_bom(data, name):
    with pytest.raises((ValueError, UnicodeError)):
        parse_bom(data, name)


def test_groups_duplicates_and_mpn():
    b = parse_bom(
        b'Value,Reference,Quantity,MPN\n1k,"R1,R2",2,PART\n1000,R3,1,PART\n2k,R1,1,PART\n,,1,\n,R4,2,\n',
        "bom.csv",
    )
    assert (
        b.components["R1"].value == "1k"
        and b.components["R1"].part_specification_id == b.components["R3"].part_specification_id
        and b.components["R4"].value is None
    )
    assert {"DUPLICATE_REFERENCE", "MISSING_REFERENCE", "QUANTITY_MISMATCH", "MISSING_VALUE"} <= {
        i.code for i in b.issues
    }


def test_formula_rejected():
    w = openpyxl.Workbook()
    w.active.append(["Reference", "Value"])
    w.active.append(["R1", "=1+1"])
    out = io.BytesIO()
    w.save(out)
    with pytest.raises(ValueError, match="Formula"):
        parse_bom(out.getvalue(), "bom.xlsx")


def test_file_limit():
    with pytest.raises(ValueError):
        parse_bom(b" " * (MAX_BYTES + 1), "bom.csv")


def test_qcv_repeated_and_conflict():
    n = QcvNetlistParser().parse(
        b"NET: 'A' R1-1 R2-1\nNET: 'A' R1-1 R3-1\nNET: 'B' R1-1\nPIN: '/U1' U1-5 (by TERM)\n", "x.qcv"
    )
    assert len(n.nets) == 2 and len(n.connections) == 2 and n.pins[-1].disposition == "TERMINATED"
    assert {i.code for i in n.issues} == {"DUPLICATE_NET", "DUPLICATE_CONNECTION", "AMBIGUOUS_CONNECTION"}


@pytest.mark.parametrize(
    "data",
    [
        b"",
        b"COMP: R1",
        b"NET: A R1-1",
        b"NET: 'A' 1R-1",
        b"PIN: '/U1' U1-1 U1-2",
        b"PIN: '/U1' U1-1 (by ???)",
        b"\xff",
    ],
)
def test_bad_qcv(data):
    with pytest.raises((ValueError, UnicodeError)):
        QcvNetlistParser().parse(data, "bad.qcv")


def test_hierarchy():
    n = QcvNetlistParser().parse(b"NET: '/VCC+3.3(V)' /sheet/U-1-A1 R1-2\n", "x.qcv")
    assert n.pins[0].key == PinKey(reference="/sheet/U-1", number="A1")
    assert next(iter(n.nets.values())).name == "/VCC+3.3(V)"


@pytest.mark.parametrize(
    "text",
    [
        "",
        "()",
        "(export",
        "(export))",
        "(export) trailing",
        '("export)',
        '(export "a)',
        "(" * 130 + "x" + ")" * 130,
    ],
)
def test_bad_sexpression(text):
    with pytest.raises(ValueError):
        sexpression(text)


def test_immutable_relations(boards):
    d = boards["fe"][2]
    with pytest.raises(TypeError):
        d.components["fake"] = d.components["U16"]
    with pytest.raises(TypeError):
        d.components["U16"].raw_sources["bom"]["Value"] = "fake"
    with pytest.raises(ValidationError):
        d.version = "fake"
    payload = wire(d)
    payload["connections"].append({"pin": {"reference": "U16", "number": "5"}, "netId": "missing"})
    with pytest.raises(ValueError):
        CircuitDesign.model_validate(payload)


def test_graph(boards):
    g = boards["fe"][3]
    p = g.get_pin_context("U16", "5")
    assert p["net"].name == "U16_EN"
    assert {x.key for x in g.get_connected_pins(p["net"].id)} == {
        PinKey(reference=r, number=n) for r, n in [("U16", "5"), ("R107", "2"), ("R83", "2")]
    }
    assert g.get_component_attributes("R107")["Resistance"].normalized == 140000
    assert g.get_component_attributes("R83")["Resistance"].normalized == 56000
    for r in ["R107", "R83"]:
        assert g.get_component(r).assembly_status == "DNP" and g.get_net_of_pin(r, "1").id != p["net"].id
    start = PinKey(reference="U16", number="5")
    assert not any(e["kind"] == "COMPONENT" for e in g.explore(start, 3, ["R107", "R83"])["edges"])
    assert any(e["kind"] == "COMPONENT" for e in g.explore(start, 3, ["R107", "R83"], True)["edges"])
    assert g.get_pin_context("U16", "unknown")["status"] == "MISSING_PIN"
    with pytest.raises(LookupError):
        g.get_component("not-present")


@pytest.mark.parametrize("kind", ["xlsx", "xls"])
def test_multiple_sheets(kind):
    out = io.BytesIO()
    if kind == "xlsx":
        w = openpyxl.Workbook()
        w.remove(w.active)
        for ref in ["R1", "R2"]:
            s = w.create_sheet(ref)
            s.append(["Reference", "Value"])
            s.append([ref, "1k"])
        w.save(out)
        w.close()
    else:
        w = xlwt.Workbook()
        for ref in ["R1", "R2"]:
            s = w.add_sheet(ref)
            s.write(0, 0, "Reference")
            s.write(0, 1, "Value")
            s.write(1, 0, ref)
            s.write(1, 1, "1k")
        w.save(out)
    assert set(parse_bom(out.getvalue(), "bom." + kind).components) == {"R1", "R2"}


def test_kicad_unconnected_library_and_ambiguity():
    n = KiCadNetlistParser().parse(
        b"""(export (libparts (libpart (lib Device)(part PART)(aliases ALIAS)(pins (pin(num 1)(name GND)(type passive))(pin(num 2)(name GND)(type passive))(pin(num 3)(name X)(type passive))))) (components(comp(ref U1)(libsource(lib Device)(part ALIAS))(sheetpath(names /hier/)(tstamps /x)))) (nets(net(code 1)(name GND)(node(ref U1)(pin 1))) (net(code 2)(name BAD)(node(ref U1)(pin 1)))))""",
        "x.net",
    )
    b = parse_bom(b"Reference,Value,DNP\nU1,IC,no\n", "bom.csv")
    g = CircuitGraph(assemble("Library", "v1", b, n))
    assert g.get_pin_context("U1", "1")["status"] == "AMBIGUOUS"
    assert g.get_pin_context("U1", "2")["status"] == "UNCONNECTED"
    assert len(g.get_pins_by_name("U1", "GND")) == 2
    assert g.get_component("U1").source_path == "/hier/"


def test_original_kicad_fixture_checksums():
    directory = Path(__file__).parent / "fixtures"
    for r in json.loads((directory / "sources.json").read_text()):
        b = (directory / r["file"]).read_bytes()
        assert len(b) == r["bytes"] and hashlib.sha256(b).hexdigest() == r["sha256"]


@pytest.mark.parametrize(
    "data",
    [
        b"NET: ' ' R1-1",
        b"NET: 'A' R1-1\nPIN: '/R1' R1-1 (by TERM)",
        b"PIN: '/U1' U1-1 (by TERM)\nNET: 'A' U1-1",
    ],
)
def test_blank_net_and_conflicting_term(data):
    if b"' '" in data:
        with pytest.raises(ValueError):
            QcvNetlistParser().parse(data, "x.qcv")
    else:
        n = QcvNetlistParser().parse(data, "x.qcv")
        assert not n.connections and any(i.code == "AMBIGUOUS_CONNECTION" for i in n.issues)
