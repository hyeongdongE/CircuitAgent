"""공개 CSV를 문자열 XLSX로 변환하고 모든 셀을 재조회하여 비교한다."""

import argparse
from pathlib import Path

import openpyxl

from .bom import csv_cells


def convert(directory: Path):
    for board in ("fe", "ntop", "fe-ntop"):
        cells = csv_cells((directory / (board + ".csv")).read_bytes())
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "BOM"
        for i, row in enumerate(cells, 1):
            for j, value in enumerate(row, 1):
                c = sheet.cell(i, j, value)
                c.data_type = "s"  # =로 시작하는 원문도 Excel 수식으로 실행되지 않는다.
        output = directory / (board + ".xlsx")
        workbook.save(output)
        workbook.close()
        check = openpyxl.load_workbook(output, read_only=False)
        try:
            for i, row in enumerate(cells, 1):
                assert [check.active.cell(i, j).value or "" for j in range(1, len(row) + 1)] == row
        finally:
            check.close()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("directory", type=Path)
    convert(parser.parse_args().directory)


if __name__ == "__main__":
    main()
