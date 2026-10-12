"""BOM CSV/XLSX/XLS 파서. 그룹 Reference를 펼치고 원문·위치·실장 상태를 보존한다."""

from __future__ import annotations

import csv
import io
import json
import re
import zipfile
from collections.abc import Mapping
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from typing import Literal

import openpyxl
import xlrd
from pydantic import Field, model_validator

from .domain import (
    AssemblyStatus,
    AttributeValue,
    BomResult,
    ComponentInstance,
    Model,
    PartSpecification,
    issue,
    wire,
)
from .units import UNITS, normalize

MAX_BYTES = 10 * 1024 * 1024
MAX_EXPANDED = 32 * 1024 * 1024
MAX_ROWS, MAX_COLUMNS, MAX_CELLS, MAX_VALUES = 50_000, 256, 250_000, 500_000
ALIASES = {
    **dict.fromkeys(
        ["reference", "references", "ref", "refs", "designator", "designators", "refdes"], "reference"
    ),
    "value": "value",
    "designation": "value",
    "qty": "quantity",
    "quantity": "quantity",
    "qtyonboard": "quantity",
    "mpn": "mpn",
    "partnumber": "mpn",
    "manufacturerpartnumber": "mpn",
    "manufacturer": "manufacturer",
    "mfr": "manufacturer",
    "type": "type",
    "componenttype": "type",
    "voltagerating": "Voltage Rating",
    "voltage": "Voltage Rating",
    "currentrating": "Current Rating",
    "current": "Current Rating",
    "resistance": "Resistance",
    "capacitance": "Capacitance",
    "inductance": "Inductance",
    "tolerance": "Tolerance",
    "nomount": "nomount",
    "dnp": "dnp",
    "dnf": "dnp",
    "donotpopulate": "dnp",
    "donotfit": "dnp",
}


def file_limit(data):
    if len(data) > MAX_BYTES:
        raise ValueError("Parser input exceeds 10 MiB")


def clean(value):
    return None if value is None or not value.strip() or value.strip() == "~" else value.strip()


def canonical(header):
    return ALIASES.get(re.sub("[^a-z0-9]", "", header.strip().lower()), header.strip())


def catalogue(headers):
    keys = {re.sub("[^a-z0-9]", "", h.strip().lower()) for h in headers}
    return {"refdes", "qtyonboard"} <= keys


def column(header, is_catalogue):
    # 내부 PartNumber와 Voltage는 제조사 MPN/정격이라고 자동 확정하지 않는다.
    if is_catalogue and header.strip().lower() == "partnumber":
        return "Part Number"
    if is_catalogue and header.strip().lower() == "voltage":
        return "Voltage"
    return canonical(header)


class BomMapping(Model):
    """자동 별칭으로 알 수 없는 헤더는 명시적으로 의미를 지정한다. 나머지 열은 모두 동적 속성이다."""

    columns: Mapping[str, str] = Field(default_factory=dict)
    units: Mapping[str, str] = Field(default_factory=dict)
    no_mount_blank: Literal["FITTED", "UNSPECIFIED"] | None = None
    encoding: Literal["utf-8-sig", "utf-8", "cp949", "euc-kr", "cp1252"] = "utf-8-sig"
    delimiter: Literal[",", ";", "\t"] | None = None

    @model_validator(mode="after")
    def checked(self):
        if len(self.columns) > MAX_COLUMNS or len(self.units) > MAX_COLUMNS:
            raise ValueError("BOM mapping column budget exceeded")
        values = list(self.columns.values())
        if len(set(values)) != len(values):
            raise ValueError("One source column cannot have conflicting roles")
        for k, v in [*self.columns.items(), *self.units.items()]:
            if not k.strip() or not v.strip() or len(k) > 512 or len(v) > 512:
                raise ValueError("Explicit bounded BOM mapping names required")
        for k, v in self.units.items():
            if v not in {"Ω", "F", "H", "V", "A", "%"}:
                raise ValueError("Unsupported explicit attribute unit")
            if k in UNITS and v != UNITS[k]:
                raise ValueError("Known physical attribute has incompatible unit")
        return self


def logical_name(name):
    # 내부 품번/일반 Voltage를 제조사 MPN/정격으로 자동 승격하지 않는다.
    return name if name in {"Part Number", "Part Name", "Voltage"} else canonical(name)


def schema(headers, config):
    explicit = {source: logical_name(target) for target, source in config.columns.items()}
    if any(h not in headers for h in explicit):
        raise ValueError("Mapped BOM source column is missing")
    mapped = {}
    claimed = set(explicit.values())
    for h in headers:
        if not h.strip():
            continue
        target = explicit.get(h, column(h, catalogue(headers)))
        if h not in explicit and target in claimed:
            target = h if h != target else "source:" + h
        if target in mapped.values():
            raise ValueError("Ambiguous duplicate BOM column; provide explicit mapping")
        mapped[h] = target
    if "reference" not in mapped.values():
        raise ValueError("Reference column or explicit reference mapping required")
    return mapped


def checked_headers(cells, config=None):
    if len([h for h in cells if h.strip()]) != len(set(h for h in cells if h.strip())):
        raise ValueError("Duplicate source BOM header")
    schema(cells, config or BomMapping())
    return cells


@dataclass(frozen=True)
class BomRow:
    location: str
    raw: dict[str, str]
    headers: tuple[str, ...]
    unnamed: tuple[str, ...]


def row_map(headers, cells):
    # 짧은 행은 빈 셀을 보존한다. 헤더 없는 추가 값은 이름을 추측하지 않고 위치 이름+Issue로 보존한다.
    raw = {h: cells[i] if i < len(cells) else "" for i, h in enumerate(headers) if h.strip()}
    unnamed = []
    for i, value in enumerate(cells):
        if value.strip() and (i >= len(headers) or not headers[i].strip()):
            key = f"column:{i + 1}"
            while key in raw or key in headers:
                key = "source:" + key
            raw[key] = value
            unnamed.append(key)
    return raw, tuple(unnamed)


def csv_cells(data, delimiter=None, encoding="utf-8-sig"):
    file_limit(data)
    text = data.decode(encoding, errors="strict")
    line = next((v for v in text.splitlines() if v.strip()), None)
    if line is None:
        raise ValueError("Empty BOM")
    quoted, counts = False, {",": 0, ";": 0, "\t": 0}
    for char in line:
        if char == '"':
            quoted = not quoted
        elif not quoted and char in counts:
            counts[char] += 1
    delimiter = delimiter or (
        "\t" if counts["\t"] > max(counts[","], counts[";"]) else ";" if counts[";"] > counts[","] else ","
    )
    # csv.reader는 quoted delimiter와 여러 줄 문자열을 표준 문법으로 읽는다.
    try:
        rows = []
        total = 0
        for row in csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, strict=True):
            if not row:
                continue
            total += len(row)
            if len(rows) >= MAX_ROWS or total > MAX_CELLS or len(row) > MAX_COLUMNS:
                raise ValueError("BOM row/column/cell limit exceeded")
            rows.append(row)
    except csv.Error as exc:
        raise ValueError("Invalid CSV content") from exc
    return rows


def validate_zip(data):
    # 압축 metadata와 실제 읽힌 bytes를 모두 검사해 XLSX ZIP bomb를 Workbook 할당 전에 차단한다.
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            entries = z.infolist()
            if len(entries) > 1000 or sum(e.file_size for e in entries) > MAX_EXPANDED:
                raise ValueError("Excel expanded content limit exceeded")
            total = 0
            for entry in entries:
                with z.open(entry) as f:
                    while chunk := f.read(8192):
                        total += len(chunk)
                        if total > MAX_EXPANDED:
                            raise ValueError("Excel expanded content limit exceeded")
    except zipfile.BadZipFile as exc:
        raise ValueError("Invalid Excel ZIP") from exc


def cell_text(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        return format(value, ".15g")
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value)


def rows_from_file(data, filename, config=None):
    config = config or BomMapping()
    file_limit(data)
    ext = filename.lower().rsplit(".", 1)[-1]
    if ext == "csv":
        rows = csv_cells(data, config.delimiter, config.encoding)
        headers = checked_headers(rows[0], config)
        result = []
        for index, row in enumerate(rows[1:], 2):
            if any(c.strip() for c in row):
                raw, unnamed = row_map(headers, row)
                result.append(BomRow(f"csv:record:{index}", raw, tuple(headers), unnamed))
        return result
    if ext not in {"xlsx", "xls"}:
        raise ValueError("Supported BOM formats: CSV, XLSX, XLS")
    expected = b"PK\x03\x04" if ext == "xlsx" else b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
    if not data.startswith(expected):
        raise ValueError("Excel content does not match filename extension")
    budget, result, found = 0, [], False
    if ext == "xlsx":
        validate_zip(data)
        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=False)
        sheets = [(s.title, s.iter_rows()) for s in workbook]
    else:
        workbook = xlrd.open_workbook(file_contents=data, on_demand=True)

        def xls_rows(sheet):
            for row in range(sheet.nrows):
                yield [sheet.cell(row, col) for col in range(sheet.ncols)]

        sheets = [(sheet.name, xls_rows(sheet)) for sheet in workbook.sheets()]
    try:
        for title, rows in sheets:
            headers = None
            for index, cells in enumerate(rows, 1):
                if index > MAX_ROWS or len(cells) > MAX_COLUMNS:
                    raise ValueError("BOM row/column limit exceeded")
                budget += len(cells)
                if budget > MAX_CELLS:
                    raise ValueError("BOM cell limit exceeded")
                # openpyxl은 POI처럼 수식을 계산하지 않는다. 수식을 값으로 잘못 해석하지 않고 명시 거절한다.
                if ext == "xlsx" and any(c.data_type == "f" for c in cells):
                    raise ValueError(
                        "Formula cells require explicit values; Python importer does not evaluate Excel formulas"
                    )
                values = [cell_text(c.value) for c in cells]
                if headers is None:
                    mapped_ref = config.columns.get("reference") or config.columns.get("Reference")
                    if any(canonical(c) == "reference" for c in values) or mapped_ref in values:
                        headers = checked_headers(values, config)
                        found = True
                    continue
                raw, unnamed = row_map(headers, values)
                if any(v.strip() for v in raw.values()):
                    result.append(BomRow(f"sheet:{title}:row:{index}", raw, tuple(headers), unnamed))
    finally:
        if ext == "xlsx":
            workbook.close()
        else:
            workbook.release_resources()
    if not found:
        raise ValueError("BOM Reference column required")
    return result


def parse_bom(data: bytes, filename: str, mapping=None) -> BomResult:
    config = mapping if isinstance(mapping, BomMapping) else BomMapping.model_validate(mapping or {})
    components, specifications, variants, issues = {}, {}, {}, []
    rows = rows_from_file(data, filename, config)
    source_headers = []
    for row in rows:
        location, raw = row.location, row.raw
        source_headers.extend(h for h in row.headers if h not in source_headers)
        resolved = schema(row.headers, config)
        fields = {resolved.get(h, h): v for h, v in raw.items()}
        references = clean(fields.get("reference"))
        if references is None:
            issues.append(issue("MISSING_REFERENCE", source=location, error=True))
            continue
        refs = [r for r in re.split(r"[,;\s]+", references) if r]
        quantity = clean(fields.get("quantity"))
        if quantity is not None:
            try:
                q = Decimal(quantity)
                if not q.is_finite() or q != q.to_integral_value():
                    raise ValueError("Quantity must be an integer")
                if q != len(refs):
                    issues.append(issue("QUANTITY_MISMATCH", references, location))
            except (ValueError, InvalidOperation):
                issues.append(issue("PARSE_FAILURE", references, location, "Invalid quantity", True))
        for ref in refs:
            if ref in components:
                issues.append(issue("DUPLICATE_REFERENCE", ref, location, "First instance retained", True))
                continue
            if row.unnamed:
                issues.append(
                    issue(
                        "PARSE_FAILURE",
                        ref,
                        location,
                        "Values without column headers retained as positional attributes: "
                        + ", ".join(row.unnamed),
                        True,
                    )
                )
            value = clean(fields.get("value"))
            if value is None:
                issues.append(issue("MISSING_VALUE", ref, location))
            type_ = clean(fields.get("type"))
            if type_ is None:
                type_ = next(
                    (
                        t
                        for prefix, t in [("R", "resistor"), ("C", "capacitor"), ("L", "inductor")]
                        if re.fullmatch(prefix + "[0-9]+", ref)
                    ),
                    None,
                )
            else:
                type_ = type_.lower()

            def placement(raw_flag, no_mount=False):
                flag = clean(raw_flag)
                if flag is None:
                    if no_mount:
                        # 확인된 공개 catalogue 규칙 또는 사용자 명시 정책만 빈 값을 실장으로 해석한다.
                        blank = config.no_mount_blank or (
                            "FITTED" if catalogue(row.headers) else "UNSPECIFIED"
                        )
                        return AssemblyStatus(blank)
                    return AssemblyStatus.UNSPECIFIED
                if no_mount:
                    result = AssemblyStatus.DNP if flag.upper() == "NM" else None
                else:
                    result = (
                        AssemblyStatus.DNP
                        if flag.lower() in {"1", "true", "yes", "y", "dnp", "dnf", "x"}
                        else AssemblyStatus.FITTED
                        if flag.lower() in {"0", "false", "no", "n"}
                        else None
                    )
                if result is None:
                    issues.append(issue("PARSE_FAILURE", ref, location, "Unknown assembly flag"))
                return result or AssemblyStatus.UNSPECIFIED

            mounted = (
                placement(fields.get("nomount"), True)
                if "nomount" in fields
                else placement(fields.get("dnp"))
            )
            if (
                "nomount" in fields
                and clean(fields.get("dnp")) is not None
                and mounted != placement(fields["dnp"])
            ):
                issues.append(
                    issue("SPECIFICATION_CONFLICT", ref, location, "NoMount and DNP disagree", True)
                )
                mounted = AssemblyStatus.UNSPECIFIED

            def attribute(key, raw_value):
                unit = config.units.get(key, UNITS.get(key))
                if unit is None:
                    return AttributeValue(raw=raw_value)
                try:
                    return normalize(raw_value, unit)
                except ValueError:
                    issues.append(issue("INVALID_UNIT", ref, location, "Invalid " + key + " unit"))
                    return AttributeValue(raw=raw_value)

            attributes = {
                k: attribute(k, v)
                for k, v in fields.items()
                if clean(v) is not None
                and k
                not in {
                    "reference",
                    "value",
                    "quantity",
                    "manufacturer",
                    "mpn",
                    "type",
                    "Id",
                    "dnp",
                    "nomount",
                }
            }
            physical = {"resistor": "Resistance", "capacitor": "Capacitance", "inductor": "Inductance"}.get(
                type_
            )
            if physical and value is not None:
                from_value = attribute(physical, fields["value"])
                explicit = attributes.setdefault(physical, from_value)
                if (
                    explicit.normalized is not None
                    and from_value.normalized is not None
                    and explicit.normalized != from_value.normalized
                ):
                    issues.append(issue("VALUE_MISMATCH", ref, location, "Explicit attribute retained"))
            mpn, manufacturer = clean(fields.get("mpn")), clean(fields.get("manufacturer"))
            spec = PartSpecification(mpn=mpn, manufacturer=manufacturer, type=type_, attributes=attributes)
            catalogue_part = clean(fields.get("Part Number"))
            identity = (
                ("mpn", manufacturer or "", mpn)
                if mpn
                else ("catalogue", catalogue_part, fields.get("Rev", ""))
                if catalogue_part
                else None
            )
            if identity:
                existing = variants.setdefault(identity, [])

                def equivalent(other):
                    if other.type != spec.type or other.attributes.keys() != spec.attributes.keys():
                        return False
                    return all(
                        a.unit == b.unit and a.normalized == b.normalized
                        if a.normalized is not None and b.normalized is not None
                        else a == b
                        for key, a in spec.attributes.items()
                        for b in [other.attributes[key]]
                    )

                shared = next((s for s in existing if equivalent(s)), None)
                if shared:
                    spec = shared
                else:
                    if existing:
                        issues.append(
                            issue(
                                "SPECIFICATION_CONFLICT",
                                ref,
                                location,
                                "Same identity has conflicting attributes",
                            )
                        )
                    existing.append(spec)
            specifications[spec.id] = spec
            components[ref] = ComponentInstance(
                reference=ref,
                value=value,
                part_specification_id=spec.id,
                raw_sources={
                    "bom": raw,
                    "provenance": {
                        "bomLocation": location,
                        "bomHeaderOrder": json.dumps(row.headers, ensure_ascii=False),
                    },
                },
                source_path=location,
                assembly_status=mounted,
            )
    return BomResult(
        components=components,
        specifications=specifications,
        issues=issues,
        source_metadata={
            "bomFormat": filename.lower().rsplit(".", 1)[-1],
            "bomEncoding": config.encoding if filename.lower().endswith(".csv") else "workbook",
            "bomRowCount": str(len(rows)),
            "bomColumns": json.dumps(source_headers, ensure_ascii=False),
            "bomMapping": json.dumps(wire(config), ensure_ascii=False, sort_keys=True),
        },
    )
